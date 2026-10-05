import csv
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from ai_cases_new import NEW, RUNTIME

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F2937")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
SEC_FILL = PatternFill("solid", fgColor="E5E7EB")
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")
thin = Side(style="thin", color="D1D5DB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

user = [tuple(r) for r in csv.reader(open("user_cases.tsv", encoding="utf-8"), delimiter="\t") if len(r) == 7]
assert len(user) == 64
all_cases = [(*c, "Yours") for c in user] + [(*c, "New") for c in NEW + RUNTIME]
ids = [c[0] for c in all_cases]
assert len(ids) == len(set(ids)), [i for i in ids if ids.count(i) > 1]

SECTIONS = OrderedDict([
    ("A. Generation - action steps", ["URL", "CMP", "NOI", "CLK", "HOV", "AMB", "CTL", "VAR", "PAR", "INT", "FIL", "SEL",
                                      "IF", "EXS", "LST", "LOOP", "WHEN", "UNT", "WAIT", "SYN"]),
    ("B. Generation - extraction (aggregation)", ["PRE", "RES", "POSV", "ARR", "DET", "NEST", "TAB", "JOIN", "AGX"]),
    ("C. Generation - inputs & robustness", ["SRC", "INS", "MIS", "QLT", "DTM"]),
    ("D. Generation - end-to-end", ["E2E"]),
    ("E. Agent at run time", ["RUN"]),
])
pref_order = [p for ps in SECTIONS.values() for p in ps]
sec_of = {p: s for s, ps in SECTIONS.items() for p in ps}
TYPE_ORDER = {"Positive": 0, "Negative": 1, "Edge": 2}

def key(c):
    p = c[0].split("-")[0]
    return (pref_order.index(p), TYPE_ORDER[c[2]], c[0])
missing = {c[0].split("-")[0] for c in all_cases} - set(pref_order)
assert not missing, missing
all_cases.sort(key=key)

wb = Workbook()

# ================================================================ Cases
ws = wb.active
ws.title = "AI Test Cases"
COLS = ["ID", "Section", "Area", "Type", "Input", "Scenario", "Expected output", "Failure signal", "Source",
        "Status", "Steps in input", "Steps generated", "Steps missed", "Extra steps", "Ran first time? (Y/N)", "Actual output / notes"]
W = [10, 24, 22, 9, 14, 44, 46, 38, 8, 10, 9, 10, 9, 9, 10, 40]
for i, (n, w) in enumerate(zip(COLS, W), 1):
    c = ws.cell(row=1, column=i, value=n)
    c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 32
for r, (cid, area, typ, inp, scen, exp, fail, src) in enumerate(all_cases, 2):
    vals = [cid, sec_of[cid.split("-")[0]], area, typ, inp, scen, exp, fail, src, "Not Run", None, None, None, None, None, None]
    for col, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=col, value=v)
        c.font, c.alignment, c.border = BODY, WRAP, BORDER
        if col in (1, 4, 9, 10, 11, 12, 13, 14, 15):
            c.alignment = CENTER
        if col >= 10:
            c.fill = INPUT_FILL
    ws.cell(row=r, column=1).font = BOLD
last = len(all_cases) + 1
ws.freeze_panes = "F2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{last}"
dv = DataValidation(type="list", formula1='"Not Run,Pass,Fail,Blocked,N/A"'); dv.add(f"J2:J{last}"); ws.add_data_validation(dv)
dvy = DataValidation(type="list", formula1='"Y,N"', allow_blank=True); dvy.add(f"O2:O{last}"); ws.add_data_validation(dvy)
dvn = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True); dvn.add(f"K2:N{last}"); ws.add_data_validation(dvn)
for val, color in [("Pass", "C6EFCE"), ("Fail", "FFC7CE"), ("Blocked", "FFEB9C"), ("N/A", "E5E7EB")]:
    ws.conditional_formatting.add(f"J2:J{last}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=color)))
for val, color in [("Positive", "15803D"), ("Negative", "B91C1C"), ("Edge", "B45309")]:
    ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(operator="equal", formula=[f'"{val}"'], font=Font(name=FONT, bold=True, color=color)))

# ================================================================ Summary
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "redflow AI Test Cases - Summary"; sm["A1"].font = Font(name=FONT, size=14, bold=True)
sm["A2"] = "Updates automatically from the AI Test Cases sheet."; sm["A2"].font = Font(name=FONT, size=10, italic=True, color="6B7280")
H = ["Section", "Area", "Cases", "Positive", "Negative", "Edge", "Pass", "Fail", "Blocked", "Not Run", "Pass % of executed"]
for i, (n, w) in enumerate(zip(H, [36, 30, 7, 9, 9, 7, 7, 7, 8, 8, 12]), 1):
    c = sm.cell(row=4, column=i, value=n)
    c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")
    sm.column_dimensions[get_column_letter(i)].width = w
A = f"'AI Test Cases'!$C$2:$C${last}"; S = f"'AI Test Cases'!$B$2:$B${last}"
T = f"'AI Test Cases'!$D$2:$D${last}"; ST = f"'AI Test Cases'!$J$2:$J${last}"
areas = OrderedDict()
for c in all_cases:
    areas.setdefault((sec_of[c[0].split("-")[0]], c[1]), None)
r = 5
for (sec, area) in areas:
    sm.cell(row=r, column=1, value=sec); sm.cell(row=r, column=2, value=area)
    sm.cell(row=r, column=3, value=f"=COUNTIFS({S},A{r},{A},B{r})")
    for i, t in enumerate(["Positive", "Negative", "Edge"]):
        sm.cell(row=r, column=4 + i, value=f'=COUNTIFS({S},A{r},{A},B{r},{T},"{t}")')
    for i, s in enumerate(["Pass", "Fail", "Blocked", "Not Run"]):
        sm.cell(row=r, column=7 + i, value=f'=COUNTIFS({S},A{r},{A},B{r},{ST},"{s}")')
    sm.cell(row=r, column=11, value=f"=IF(G{r}+H{r}=0,0,G{r}/(G{r}+H{r}))").number_format = "0%"
    for col in range(1, 12):
        sm.cell(row=r, column=col).font = BODY; sm.cell(row=r, column=col).border = BORDER
    r += 1
tot = r
sm.cell(row=tot, column=1, value="TOTAL")
for col in range(3, 11):
    L = get_column_letter(col)
    sm.cell(row=tot, column=col, value=f"=SUM({L}5:{L}{tot-1})")
sm.cell(row=tot, column=11, value=f"=IF(G{tot}+H{tot}=0,0,G{tot}/(G{tot}+H{tot}))").number_format = "0%"
for col in range(1, 12):
    c = sm.cell(row=tot, column=col); c.font, c.border, c.fill = BOLD, BORDER, SEC_FILL

m = tot + 2
sm.cell(row=m, column=1, value="Generation accuracy (from filled-in counts)").font = Font(name=FONT, size=12, bold=True)
CS = f"'AI Test Cases'!$K$2:$K${last}"; CG = f"'AI Test Cases'!$L$2:$L${last}"
CM = f"'AI Test Cases'!$M$2:$M${last}"; CX = f"'AI Test Cases'!$N$2:$N${last}"; CY = f"'AI Test Cases'!$O$2:$O${last}"
metrics = [
    ("Steps in inputs (total)", f"=SUM({CS})", "0"),
    ("Steps generated (total)", f"=SUM({CG})", "0"),
    ("Steps missed (total)", f"=SUM({CM})", "0"),
    ("Extra steps (total)", f"=SUM({CX})", "0"),
    ("Recall: steps captured / steps in input", f"=IF(B{m+1}=0,0,(B{m+1}-B{m+3})/B{m+1})", "0.0%"),
    ("Precision: correct steps / steps generated", f"=IF(B{m+2}=0,0,(B{m+2}-B{m+4})/B{m+2})", "0.0%"),
    ("First-run success rate (Y / answered)", f'=IF(COUNTIF({CY},"Y")+COUNTIF({CY},"N")=0,0,COUNTIF({CY},"Y")/(COUNTIF({CY},"Y")+COUNTIF({CY},"N")))', "0.0%"),
]
for i, (lab, f, fmt) in enumerate(metrics, 1):
    a = sm.cell(row=m + i, column=1, value=lab); a.font, a.border = BODY, BORDER
    b = sm.cell(row=m + i, column=2, value=f); b.font, b.border = BOLD, BORDER; b.number_format = fmt
sm.freeze_panes = "C5"

# ================================================================ How to Use
hu = wb.create_sheet("How to Use", 0)
hu.column_dimensions["A"].width = 24; hu.column_dimensions["B"].width = 100
rows = [
    ("redflow AI Test Cases", None),
    (None, None),
    ("What this covers", "Only the AI side: (1) turning a video / PDF / screenshots into a redflow, and (2) the Agent following that redflow on a live page. UI pages (Agent list, settings, uploads) are out of scope."),
    ("Your cases", "Your 64 cases are kept unchanged (Source = Yours). New cases fill gaps: URL line, missing steps, noise, ambiguous/small/large elements, other controls, INTENT, syntax, parameters, input types, instructions, input quality, determinism, export flows, and runtime behaviour (Source = New)."),
    ("How to run a generation case", "Prepare the input described in Scenario, generate, then compare the redflow with Expected output. It fails if anything in Failure signal appears."),
    ("How to run a runtime case (RUN-*)", "Execute the redflow on the live app and check Body Cam and the app state against Expected output."),
    ("Fill in (yellow)", "Status, and for generation cases the step counts (Steps in input, Steps generated, Steps missed, Extra steps) plus Ran first time? (Y/N). Summary turns these into recall, precision and first-run success rate."),
    ("Example", "CMP-P1 | Status: Fail | Steps in input 12 | Generated 11 | Missed 1 | Extra 0 | Ran first time N | Notes: 'Next on page 2 of wizard missing'"),
    ("Type", "Positive = the feature should be used. Negative = it should NOT be used (checks over-generation). Edge = tricky variant."),
]
for i, (a, b) in enumerate(rows, 1):
    if a: hu.cell(row=i, column=1, value=a).font = Font(name=FONT, size=14, bold=True) if i == 1 else BOLD
    if b:
        c = hu.cell(row=i, column=2, value=b); c.font, c.alignment = BODY, WRAP

wb.save("Redblock_redflow_AI_Test_Cases.xlsx")
print(len(all_cases), "cases;", len(user), "yours;", len(NEW) + len(RUNTIME), "new;", "areas", len(areas))
